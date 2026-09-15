#!/usr/bin/env python3
"""Convierte una presentación (PPTX o PDF) en slides/slides.json + imágenes.

Uso:
    python3 tools/build_slides.py presentacion.pptx [--title "Título"] [--out slides]
        [--groups "1;2-4;5-6;7-17;..."] [--brief "1,39"] [--footer "texto de pie"]

--groups  agrupa páginas consecutivas (animaciones por pasos) en una sola diapositiva:
          el avatar la comenta una vez mientras se muestran los pasos.
--brief   diapositivas (tras agrupar) que el avatar debe tratar muy brevemente (índices, portadas).
--footer  texto repetido en el pie de página que se elimina del contenido.

Requisitos: LibreOffice (soffice) para PPTX, poppler-utils (pdftoppm, pdftotext),
y `pip install python-pptx` para extraer títulos, viñetas y notas del orador.
Las notas del orador se envían a la IA como guía de lo que conviene subrayar.
"""
import argparse, json, re, shutil, subprocess, sys, tempfile
from pathlib import Path


def run(cmd):
    subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def pptx_text(path):
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for s in prs.slides:
        title, bullets, other = "", [], []
        if s.shapes.title is not None and s.shapes.title.has_text_frame:
            title = s.shapes.title.text_frame.text.strip()
        for sh in s.shapes:
            if sh == s.shapes.title:
                continue
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        bullets.append(t)
            elif getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    other.append(" | ".join(c.text.strip() for c in row.cells))
        notes = ""
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            notes = s.notes_slide.notes_text_frame.text.strip()
        out.append({"title": title, "bullets": bullets[:25], "text": "\n".join(other)[:3000], "notes": notes[:4000]})
    return out, (prs.core_properties.title or "")


def pdf_text(pdf, n):
    out = []
    for i in range(1, n + 1):
        t = subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), str(pdf), "-"],
                           capture_output=True, text=True).stdout
        lines = [re.sub(r"\s+", " ", l).strip() for l in t.splitlines() if l.strip()]
        # título = primera línea con contenido textual (evita cifras sueltas de gráficos)
        k = next((j for j, l in enumerate(lines) if sum(c.isalpha() for c in l) >= 5), 0)
        lines = lines[k:k + 1] + lines[:k] + lines[k + 1:]
        out.append({"title": lines[0][:160] if lines else "", "bullets": [], "text": "\n".join(lines[1:])[:3000], "notes": ""})
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("source")
    ap.add_argument("--title", default="")
    ap.add_argument("--out", default="slides")
    ap.add_argument("--dpi", type=int, default=96)
    ap.add_argument("--groups", default="")
    ap.add_argument("--brief", default="")
    ap.add_argument("--footer", default="")
    a = ap.parse_args()

    src = Path(a.source).resolve()
    out = Path(a.out); img_dir = out / "img"
    if img_dir.exists():
        shutil.rmtree(img_dir)
    img_dir.mkdir(parents=True)

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        texts, meta_title = [], ""
        if src.suffix.lower() in (".pptx", ".ppt"):
            texts, meta_title = pptx_text(src) if src.suffix.lower() == ".pptx" else ([], "")
            run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(src)])
            pdf = tmp / (src.stem + ".pdf")
        elif src.suffix.lower() == ".pdf":
            pdf = src
        else:
            sys.exit("Formato no soportado: usa .pptx o .pdf")

        run(["pdftoppm", "-jpeg", "-jpegopt", "quality=78", "-r", str(a.dpi), str(pdf), str(tmp / "p")])
        pages = sorted(tmp.glob("p-*.jpg"), key=lambda p: int(p.stem.split("-")[-1]))
        if not texts:
            texts = pdf_text(pdf, len(pages))
        n = len(pages)
        if a.footer:
            for t in texts:
                for k in ("title", "text"):
                    t[k] = "\n".join(l for l in t.get(k, "").splitlines() if a.footer not in l)
                while not sum(c.isalpha() for c in t["title"]) >= 5 and t.get("text"):
                    first, _, rest = t["text"].partition("\n"); t["title"], t["text"] = first, rest
                if sum(c.isalpha() for c in t["title"]) < 5:
                    t["title"] = ""

        groups = []
        if a.groups:
            for part in a.groups.split(";"):
                lo, _, hi = part.strip().partition("-")
                groups.append(list(range(int(lo), int(hi or lo) + 1)))
            used = sorted(p for g in groups for p in g)
            if used != list(range(1, n + 1)):
                sys.exit(f"--groups debe cubrir las páginas 1..{n} exactamente una vez")
        else:
            groups = [[i] for i in range(1, n + 1)]
        brief = {int(x) for x in a.brief.split(",") if x.strip()}

        for i, page in enumerate(pages):
            shutil.copy(page, img_dir / f"p-{i + 1:03d}.jpg")
        slides = []
        for gi, g in enumerate(groups, 1):
            t = texts[g[-1] - 1] if g[-1] - 1 < len(texts) else {}
            s = {"pages": [g[0], g[-1]], "images": [f"{out.name}/img/p-{p:03d}.jpg" for p in g]}
            s.update({k: v for k, v in t.items() if v})
            if gi in brief:
                s["brief"] = True
            slides.append(s)

    from PIL import Image
    w, h = Image.open(img_dir / "p-001.jpg").size
    aspect = f"{w} / {h}"
    title = a.title or meta_title or (slides[0].get("title") if slides else "") or src.stem
    (out / "slides.json").write_text(json.dumps({"title": title, "source": src.name, "aspect": aspect, "slides": slides},
                                                ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{len(slides)} diapositivas → {out/'slides.json'}")


if __name__ == "__main__":
    main()
